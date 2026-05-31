from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "report-images"
REPORT = ROOT / "实验11_121072021030_林立洲.docx"


def set_font(run, font_name="宋体", size=11):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)


def add_heading(document, text, level=1):
    paragraph = document.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_font(run, "黑体", 16 if level == 1 else 14)
    run.bold = True
    return paragraph


def add_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(22)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_plain_paragraph(document, text, align=None):
    paragraph = document.add_paragraph()
    if align:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(code.strip())
    set_font(run, "Consolas", 9)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph._p.get_or_add_pPr().append(shading)
    return paragraph


def add_picture(document, filename, caption, width=6.3):
    path = IMG / filename
    if not path.exists():
        add_plain_paragraph(document, f"图缺失：{filename}")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_font(caption_run, "宋体", 10)


def add_info_table(document):
    table = document.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = [
        ("实验名称", "实验11 状态模式"),
        ("课程名称", "软件体系结构与设计模式"),
        ("年级专业", "2023级软件工程"),
        ("班级", "软工2班"),
        ("学号", "121072021030"),
        ("姓名", "林立洲"),
        ("实验环境", "Windows、Enterprise Architect 12、IntelliJ IDEA、Java、Maven"),
        ("项目名称", "experiment11-state-observer"),
    ]
    for row, (key, value) in zip(table.rows, rows):
        row.cells[0].text = key
        row.cells[1].text = value
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("《软件体系结构与设计模式》实验报告")
    set_font(run, "黑体", 20)
    run.bold = True

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("实验11 状态模式与课堂观察者模式任务")
    set_font(run, "黑体", 16)
    run.bold = True

    add_info_table(document)

    add_heading(document, "一、实验目的")
    add_paragraph(document, "本次实验通过观察者模式和状态模式两个任务，训练先分析业务需求、再进行 UML 建模、最后使用 Java 编码实现并验证的完整开发过程。观察者模式部分主要掌握目标对象与观察者对象之间的一对多通知机制；状态模式部分主要掌握对象状态变化导致行为变化的设计思想。")
    add_paragraph(document, "通过机房温度监控场景，可以理解当一个目标对象发生变化时，多个依赖对象如何自动收到通知并执行响应动作。通过 WMS 仓储系统采购单场景，可以理解在业务流程较长、状态规则较多时，如何把不同状态下的操作权限封装到独立状态类中，从而避免在环境类中堆积大量条件判断。")
    add_paragraph(document, "实验还要求使用 Enterprise Architect 绘制 UML 类图、状态图和顺序图，并将设计结果与 Java 代码实现保持一致。通过 Maven 构建、主程序运行、非法操作拦截和 Git 仓库整理，验证系统设计是否符合课堂任务和实验要求。")
    add_paragraph(document, "在完成基础要求的基础上，本项目进一步实现了实验文档中给出的加分项和进阶功能，包括责任链校验、管理员权限控制、部分入库、部分开票以及工厂模式创建入库单、发票、退货单，使程序更接近真实 WMS 采购业务。")

    add_heading(document, "二、实验任务与需求分析")
    add_heading(document, "2.1 课堂任务：观察者模式", 2)
    add_paragraph(document, "课堂任务要求设计一个机房监控系统。当温度传感器检测到机房温度达到指定阈值时，需要把高温信号传递给多个响应设备。响应设备包括警示灯、报警器、安全逃生门和隔热门。警示灯负责闪烁，报警器负责报警，安全逃生门自动开启，隔热门自动关闭。")
    add_paragraph(document, "该场景中温度传感器与响应设备之间是一对多关系，传感器不应该直接绑定某一个具体设备，而是应当维护观察者集合。这样当后续增加新的响应设备时，只需要实现统一的观察者接口并注册到传感器中，不需要修改传感器核心逻辑。因此本任务适合采用观察者模式。")

    add_heading(document, "2.2 实验11：状态模式", 2)
    add_paragraph(document, "实验11要求基于状态模式实现 WMS 仓储系统采购单的完整生命周期。采购单状态包括草稿、已审批、已入库、已开票、已付款、已取消和已退货完结。系统需要支持编辑采购单、审批采购单、取消采购单、生成入库单、生成发票、退货退款和付款七类业务操作。")
    add_paragraph(document, "采购单业务规则的核心特点是状态决定行为。例如草稿状态允许编辑、审批和取消，但不允许直接入库、开票或付款；已审批状态允许生成入库单和取消，但不允许重复审批；已入库状态允许生成发票或退货退款，但禁止直接取消；已开票状态允许付款或退货退款；已取消和已退货完结属于终止状态，后续业务操作均应被拦截。")
    add_paragraph(document, "如果把所有状态规则都写在 PurchaseOrder 中，代码会逐渐变成大量 if-else 或 switch 判断，既不利于阅读，也不利于扩展。本实验采用状态模式，将每种状态封装为独立状态类，由当前状态对象负责处理业务操作和状态流转。")
    add_paragraph(document, "实验文档还提出了进阶加分方向：可以搭配责任链或策略模式优化校验结构，增加日志记录和权限校验，实现部分入库、部分开票场景，并结合工厂模式统一创建入库单、发票和退货单。本项目对这些内容都进行了实现和测试。")

    add_heading(document, "三、系统设计")
    add_heading(document, "3.1 观察者模式设计", 2)
    add_paragraph(document, "观察者模式设计中，Subject 表示抽象目标，定义 attach、detach 和 notifyObservers 方法；AlarmObserver 表示抽象观察者，定义 update 方法；TemperatureSensor 是具体目标对象，保存当前温度、报警阈值和观察者集合；WarningLight、AlarmBell、EscapeDoor、InsulationDoor 是具体观察者，分别完成警示灯闪烁、报警器报警、逃生门开启和隔热门关闭。")
    add_paragraph(document, "ObserverTaskDemo 作为客户端测试类，负责创建温度传感器和四类响应设备，将响应设备注册到传感器中，并模拟温度达到 80.0℃ 的高温场景。传感器在检测到温度超过阈值后，统一调用 notifyObservers 通知所有观察者。")
    add_picture(document, "01_observer_class_diagram.png", "图1 课堂任务机房监控观察者模式类图")
    add_picture(document, "02_observer_sequence_diagram.png", "图2 课堂任务机房监控观察者模式顺序图", width=6.7)

    add_heading(document, "3.2 状态模式设计", 2)
    add_paragraph(document, "状态模式设计中，PurchaseOrder 是环境类 Context，内部持有当前状态对象 state、操作日志集合 operationLogs 和业务单据集合 businessDocuments。PurchaseOrderState 是抽象状态接口，规定采购单所有业务操作的方法。AbstractPurchaseOrderState 是抽象状态基类，为非法操作提供统一默认处理，默认抛出 BusinessException 并说明失败原因。")
    add_paragraph(document, "DraftState、ApprovedState、PartiallyInStockState、InStockState、PartiallyInvoicedState、InvoicedState、PaidState、CancelledState、ReturnedState 分别表示采购单的不同业务状态。PartiallyInStockState 用于表示采购单已经生成过部分入库单但尚未全部入库，PartiallyInvoicedState 用于表示采购单已经部分开票但尚未满足付款条件。")
    add_paragraph(document, "权限设计中使用 UserRole 区分普通操作员和管理员。普通操作员不能撤销已审批采购单，管理员可以在已审批状态执行 revokeApproval，将采购单退回草稿状态重新编辑和审批。该功能对应实验拓展要求中的权限简易校验。")
    add_paragraph(document, "业务单据设计中使用 BusinessDocumentFactory 统一创建 InboundOrderDocument、InvoiceDocument 和 ReturnOrderDocument。每次成功入库、开票或退货退款时，采购单都会生成对应业务单据并保存到 businessDocuments 集合中，最终在 showInfo 中一起输出。")
    add_paragraph(document, "公共校验设计中使用责任链模式。state.validation 包中的 OrderValidationRule 负责串联不同校验节点，OrderValidationChains 统一组织编辑校验、业务数量校验和管理员权限校验。这样公共校验逻辑不需要散落在各个状态类中，更符合单一职责和开闭原则。")
    add_paragraph(document, "本实验将已付款状态视为付款完成状态，不再从已付款状态触发退货退款。退货退款主要在已入库、部分开票和已开票状态触发，这样可以覆盖退货退款功能，同时避免与付款完结逻辑冲突。OperationLog 用于记录每次合法操作的操作名称、原状态、新状态、操作结果和时间，便于测试时查看状态流转历史。")
    add_picture(document, "03_state_class_diagram.png", "图3 实验11 WMS采购单状态模式类图", width=6.7)
    add_picture(document, "04_state_state_diagram.png", "图4 实验11 WMS采购单状态流转图", width=6.7)
    add_picture(document, "05_state_normal_sequence.png", "图5 实验11 WMS采购单正常流程顺序图", width=6.7)
    add_picture(document, "06_state_invalid_sequence.png", "图6 实验11 WMS采购单非法操作拦截顺序图", width=6.7)
    add_picture(document, "07_state_return_sequence.png", "图7 实验11 WMS采购单退货退款流程顺序图", width=6.7)

    add_heading(document, "四、系统实现")
    add_heading(document, "4.1 项目结构", 2)
    add_paragraph(document, "项目使用 Maven 管理，主入口为 com.lzl.experiment11.MainApp。observer 包存放课堂观察者模式任务代码，state 包存放实验11状态模式代码，state.validation 包存放责任链校验代码。报告图片、运行结果和 EA 项目文件保存在项目根目录下，便于统一提交和检查。")
    add_picture(document, "08_project_structure.png", "图8 项目结构截图", width=5.8)

    add_heading(document, "4.2 观察者模式核心代码", 2)
    add_paragraph(document, "AlarmObserver 是所有响应设备共同实现的接口，TemperatureSensor 通过 Subject 接口维护观察者集合。传感器只依赖 AlarmObserver 抽象接口，不直接依赖具体设备类。")
    add_code(document, """
public interface AlarmObserver {
    void update(double temperature);
}

public class TemperatureSensor implements Subject {
    private double temperature;
    private final double threshold;
    private final List<AlarmObserver> observers = new ArrayList<>();

    public void attach(AlarmObserver observer) {
        observers.add(observer);
    }

    public void notifyObservers() {
        for (AlarmObserver observer : observers) {
            observer.update(temperature);
        }
    }
}
""")
    add_paragraph(document, "当 setTemperature 方法接收到新的温度值后，调用 checkTemperature 判断是否超过阈值。若超过阈值，传感器不需要知道具体响应设备是什么，只需要统一调用观察者的 update 方法。")
    add_code(document, """
public void setTemperature(double temperature) {
    this.temperature = temperature;
    System.out.printf("温度传感器检测到当前温度：%.1f℃%n", temperature);
    checkTemperature();
}

public void checkTemperature() {
    if (temperature >= threshold) {
        notifyObservers();
        return;
    }
    System.out.printf("当前温度未达到报警阈值 %.1f℃，系统保持正常监控。%n", threshold);
}
""")
    add_paragraph(document, "具体观察者只负责自己的响应动作。例如 WarningLight 在收到高温通知后调用 flash 方法，实现警示灯闪烁。")
    add_code(document, """
public class WarningLight implements AlarmObserver {
    public void update(double temperature) {
        System.out.printf("警示灯接收到 %.1f℃ 高温信号。%n", temperature);
        flash();
    }

    public void flash() {
        System.out.println("警示灯开始闪烁");
    }
}
""")

    add_heading(document, "4.3 状态模式核心代码", 2)
    add_paragraph(document, "PurchaseOrder 的业务方法只负责把请求委托给当前状态对象，不在环境类中集中判断状态。这样新增状态或调整状态规则时，修改范围主要集中在具体状态类。")
    add_code(document, """
public void approve() {
    state.approve(this);
}

public void generateInboundOrder() {
    generateInboundOrder(getRemainingInboundQuantity());
}

public void generateInboundOrder(int inboundQuantity) {
    state.generateInboundOrder(this, inboundQuantity);
}

public void pay() {
    state.pay(this);
}
""")
    add_paragraph(document, "PurchaseOrderState 接口统一规定采购单支持的业务操作，所有具体状态类都通过该接口对外提供一致的处理方式。")
    add_code(document, """
public interface PurchaseOrderState {
    void edit(PurchaseOrder order, String productName, int quantity, String supplier);
    void approve(PurchaseOrder order);
    void cancel(PurchaseOrder order);
    void revokeApproval(PurchaseOrder order, UserRole role);
    void generateInboundOrder(PurchaseOrder order, int inboundQuantity);
    void generateInvoice(PurchaseOrder order, int invoiceQuantity);
    void pay(PurchaseOrder order);
    void returnAndRefund(PurchaseOrder order);
    String getStateName();
}
""")
    add_paragraph(document, "AbstractPurchaseOrderState 为非法操作提供默认实现。具体状态类只需要覆盖自己允许的操作，未覆盖的操作自动进入统一拦截逻辑。")
    add_code(document, """
protected void reject(String operation, String reason) {
    throw new BusinessException(
        operation + "失败：当前状态【" + getStateName() + "】不满足业务规则。" + reason
    );
}
""")
    add_paragraph(document, "DraftState 允许编辑、审批和取消。审批成功后切换到 ApprovedState，取消成功后切换到 CancelledState。")
    add_code(document, """
public class DraftState extends AbstractPurchaseOrderState {
    public void approve(PurchaseOrder order) {
        order.changeState(new ApprovedState(), "审批采购单",
                "审批通过，采购单进入已审批状态。");
    }

    public void cancel(PurchaseOrder order) {
        order.changeState(new CancelledState(), "取消采购单",
                "采购单尚未入库，取消成功。");
    }
}
""")
    add_paragraph(document, "ApprovedState 允许生成入库单、取消采购单和管理员撤销审批。生成入库单时可以一次全量入库，也可以先进入部分入库状态，后续由 PartiallyInStockState 继续补齐入库。")
    add_code(document, """
public class ApprovedState extends AbstractPurchaseOrderState {
    public void revokeApproval(PurchaseOrder order, UserRole role) {
        OrderValidationChains.validateAdmin("撤销审批", role);
        order.changeState(new DraftState(), "撤销审批",
                role.getDisplayName() + "撤销已审批采购单，单据退回草稿状态。");
    }

    public void generateInboundOrder(PurchaseOrder order, int inboundQuantity) {
        OrderValidationChains.validateBusinessQuantity(
                "生成入库单", inboundQuantity, order.getRemainingInboundQuantity());
        BusinessDocument document = order.recordInboundDocument(inboundQuantity);
        if (order.isInboundComplete()) {
            order.changeState(new InStockState(), "生成入库单",
                    document.getDocumentNo() + " 已生成，采购单全部入库。");
            return;
        }
        order.changeState(new PartiallyInStockState(), "部分入库",
                document.getDocumentNo() + " 已生成，当前已入库 "
                        + order.getInboundQuantity() + "/" + order.getQuantity() + "。");
    }
}
""")
    add_paragraph(document, "InStockState 允许生成发票和退货退款；PartiallyInvoicedState 表示已经部分开票但尚未满足付款条件，继续补齐发票后才会进入 InvoicedState。")
    add_code(document, """
public class InStockState extends AbstractPurchaseOrderState {
    public void generateInvoice(PurchaseOrder order, int invoiceQuantity) {
        OrderValidationChains.validateBusinessQuantity(
                "生成发票", invoiceQuantity, order.getRemainingInvoiceQuantity());
        BusinessDocument document = order.recordInvoiceDocument(invoiceQuantity);
        if (order.isInvoiceComplete()) {
            order.changeState(new InvoicedState(), "生成发票",
                    document.getDocumentNo() + " 已生成，采购单全部开票。");
            return;
        }
        order.changeState(new PartiallyInvoicedState(), "部分开票",
                document.getDocumentNo() + " 已生成，当前已开票 "
                        + order.getInvoiceQuantity() + "/" + order.getQuantity() + "。");
    }

    public void returnAndRefund(PurchaseOrder order) {
        BusinessDocument document = order.recordReturnDocument(order.getInboundQuantity());
        order.changeState(new ReturnedState(), "退货退款",
                document.getDocumentNo() + " 已生成，退货退款成功，流程完结。");
    }
}
""")
    add_paragraph(document, "业务单据工厂负责根据单据类型创建不同业务单据，PurchaseOrder 不需要直接关心具体单据类的创建细节。")
    add_code(document, """
public final class BusinessDocumentFactory {
    private static final Map<DocumentType, DocumentCreator> CREATORS =
            new EnumMap<>(DocumentType.class);

    static {
        CREATORS.put(DocumentType.INBOUND_ORDER, InboundOrderDocument::new);
        CREATORS.put(DocumentType.INVOICE, InvoiceDocument::new);
        CREATORS.put(DocumentType.RETURN_ORDER, ReturnOrderDocument::new);
    }

    public static BusinessDocument create(DocumentType type, PurchaseOrder order, int quantity) {
        return CREATORS.get(type).create(order, quantity);
    }
}
""")
    add_paragraph(document, "责任链校验把商品名称、数量、供应商和管理员权限等公共校验拆成独立规则，再由 OrderValidationChains 按场景组装。")
    add_code(document, """
public static void validateEdit(String productName, int quantity, String supplier) {
    OrderValidationRule chain = new ProductNameRequiredRule();
    chain.linkWith(new PositiveQuantityRule())
            .linkWith(new SupplierRequiredRule());
    chain.check(ValidationContext.edit(productName, quantity, supplier));
}

public static void validateAdmin(String operation, UserRole role) {
    OrderValidationRule chain = new AdminRoleRule();
    chain.check(ValidationContext.admin(operation, role));
}
""")
    add_paragraph(document, "每次合法操作都会通过 recordOperation 或 changeState 写入 OperationLog，并在控制台输出操作名称、原状态、新状态和操作结果；每次成功生成入库单、发票或退货单时也会保存业务单据，便于测试时追踪完整生命周期。")

    add_heading(document, "五、测试与运行结果")
    add_heading(document, "5.1 Maven 构建结果", 2)
    add_paragraph(document, "项目使用 Maven 执行 clean package，编译 40 个 Java 源文件并生成 target/experiment11-state-observer-1.0.0.jar。构建结果显示 BUILD SUCCESS。")
    add_picture(document, "14_maven_package_success.png", "图9 Maven package 成功结果", width=6.5)

    add_heading(document, "5.2 观察者模式运行结果", 2)
    add_paragraph(document, "运行课堂任务时，程序先注册 WarningLight、AlarmBell、EscapeDoor 和 InsulationDoor 四个响应设备。随后传感器检测到当前温度为 80.0℃，超过 60.0℃ 阈值，因此通知所有观察者。运行结果表明四个设备都按设计完成了响应。")
    add_picture(document, "10_observer_run_result.png", "图10 观察者模式运行结果", width=6.2)

    add_heading(document, "5.3 状态模式正常流程测试", 2)
    add_paragraph(document, "正常付款流程使用采购单 PO-001。该采购单依次执行编辑、审批、生成入库单、生成发票和付款操作，状态从草稿流转到已审批、已入库、已开票，最终到达已付款。运行结果中可以看到系统生成了入库单和发票，并记录了完整操作日志。")
    add_picture(document, "11_state_normal_result.png", "图11 状态模式正常流程运行结果", width=6.5)

    add_heading(document, "5.4 状态模式取消流程测试", 2)
    add_paragraph(document, "取消流程使用采购单 PO-002。采购单先在草稿状态完成编辑，再审批为已审批状态。由于此时还没有入库，系统允许取消采购单，最终状态为已取消。该流程验证了入库前允许取消、入库后不允许取消的业务边界。")

    add_heading(document, "5.5 状态模式退货退款流程测试", 2)
    add_paragraph(document, "退货退款流程使用采购单 PO-003。该采购单完成编辑、审批和入库后，处于已入库状态。在该状态下调用 returnAndRefund 方法，由 InStockState 生成退货单并将状态切换为 ReturnedState，最终状态为已退货完结。")
    add_picture(document, "13_state_return_result.png", "图12 退货退款流程运行结果", width=6.5)

    add_heading(document, "5.6 非法操作拦截测试", 2)
    add_paragraph(document, "非法流程测试覆盖草稿状态直接生成入库单、草稿状态直接付款、已入库状态取消采购单、已开票后再次编辑采购单、已取消后再次审批、编辑时数量为 0 等场景。所有非法操作均由当前状态类、抽象状态基类或责任链校验规则抛出 BusinessException，并输出清晰的中文失败原因。")
    add_picture(document, "12_state_invalid_result.png", "图13 非法操作拦截运行结果", width=6.5)
    add_heading(document, "5.7 进阶加分功能测试", 2)
    add_paragraph(document, "进阶流程测试覆盖两个重点场景。PO-BONUS-001 先由普通操作员尝试撤销已审批采购单，系统通过权限责任链拦截；随后管理员撤销审批成功，采购单退回草稿并重新审批。PO-BONUS-002 演示部分入库、补齐入库、部分开票、补齐发票和付款，验证部分流程不会造成状态死锁。")
    add_paragraph(document, "该流程还验证了业务单据工厂的效果：部分入库时分别生成两张入库单，部分开票时分别生成两张发票，最终采购单信息中能够看到所有关联单据和操作日志。")
    add_picture(document, "16_advanced_bonus_result.png", "图14 进阶加分功能运行结果", width=6.5)
    add_paragraph(document, "完整运行结果已保存到 run-result.txt，能够看到 MainApp 按顺序执行观察者模式、状态模式基础流程和进阶加分流程的全部测试。")
    add_picture(document, "09_mainapp_run_result.png", "图15 MainApp 完整运行结果截图", width=6.5)

    add_heading(document, "六、Git 仓库说明")
    add_paragraph(document, "本项目已在本地建立 Git 仓库，并按要求提交本次实验代码、EA 项目文件、UML 图片、运行结果文件、README 和实验报告。项目已推送到 GitHub 公开仓库，仓库地址为：https://github.com/heixuan-vsmenghu/experiment11-state-observer-121072021030。")
    add_paragraph(document, "后续如果继续修改项目，可使用以下命令将本地提交推送到远程仓库：")
    add_code(document, """
git remote add origin <仓库地址>
git branch -M main
git push -u origin main
""")
    if (IMG / "15_git_commit_success.png").exists():
        add_picture(document, "15_git_commit_success.png", "图16 Git commit 成功结果", width=6.5)

    add_heading(document, "七、实验总结")
    add_paragraph(document, "通过本次实验，我进一步理解了观察者模式中一对多通知机制的设计思想。温度传感器只负责维护观察者集合并发布通知，具体响应设备只负责各自的响应动作。这种方式降低了传感器与响应设备之间的耦合，也使后续扩展新的响应设备更加方便。")
    add_paragraph(document, "通过状态模式实现 WMS 采购单生命周期，我体会到“状态决定行为”的业务建模方式。采购单在不同状态下允许执行的操作不同，如果把所有判断都写在 PurchaseOrder 中，代码会很快变得复杂。本实验将状态规则分散到不同状态类中后，正常流程、取消流程、退货退款流程、部分入库、部分开票和非法操作拦截都更清晰。")
    add_paragraph(document, "在进阶功能实现中，我进一步体会到多种设计模式之间可以配合使用：状态模式负责业务生命周期，责任链模式负责公共校验，工厂模式负责统一创建业务单据。这样每一类变化都有相对独立的扩展位置，比单纯堆叠条件判断更容易维护。")
    add_paragraph(document, "本次实验中较难的部分是把业务规则转换成状态流转关系，尤其是已入库、已开票、已付款、已取消和已退货完结之间的边界。通过先绘制 UML 状态图和顺序图，再编写代码，可以在实现前明确每一步状态变化，减少编码时的混乱。")
    add_paragraph(document, "总体来看，观察者模式适合处理对象之间的通知关系，状态模式适合处理状态较多、行为随状态变化的业务流程。通过本次实验，我不仅完成了 Java 编码，也完成了 UML 建模、运行验证、日志输出和 Git 仓库整理，形成了比较完整的实验交付过程。")

    add_heading(document, "八、附录")
    add_paragraph(document, "运行命令：mvn clean package；java -Dfile.encoding=UTF-8 -Dsun.stdout.encoding=UTF-8 -Dstdout.encoding=UTF-8 -jar target/experiment11-state-observer-1.0.0.jar。")
    add_paragraph(document, "重要文件包括 pom.xml、README.md、src/main/java/com/lzl/experiment11、实验11_121072021030_林立洲.eap、report-images、run-result.txt、maven-package-result.txt、git-result.txt 和本 Word 报告。")
    add_paragraph(document, "重要截图清单包括 7 张 EA UML 图、项目结构截图、MainApp 完整运行结果截图、观察者模式运行结果截图、状态模式正常流程截图、非法操作拦截图、退货退款流程截图、进阶加分流程截图、Maven 构建截图和 Git 提交截图。")

    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
